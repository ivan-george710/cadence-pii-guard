"""Synthetic default-engine and live-HTTP checks; never opens real user documents.

Run engines, api, restart run.py, then reload with the same --work-dir.
Evidence and generated documents stay under .private_review_app. Requires dev deps.
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import httpx
import pymupdf as fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

from app.engines.audit import AuditEngine
from app.engines.audit.storage import AuditStore
from app.pipeline import DocumentPipeline
from tools.create_phase11_fixtures import create as create_faces

LINES = ["Owner: John Smith.", "Employee ID: EMP-00001.", "Contact: ada@example.com."]
HEADERS = {"X-Cadence-Local": "1"}


def write_json(path, data):
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def fixtures(work):
    marker = work / "fixtures.json"
    if marker.exists():
        return {name: Path(path) for name, path in json.loads(marker.read_text()).items()}
    target = work / "fixtures"
    if target.exists():
        raise RuntimeError("Refusing to overwrite an existing fixture directory")
    target.mkdir()
    native = target / "text.pdf"
    pdf = Canvas(str(native))
    for n, line in enumerate(LINES):
        pdf.drawString(60, 730 - n * 35, line)
    pdf.save()
    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 40)
    for n, line in enumerate(LINES):
        draw.text((150, 200 + n * 90), line, font=font, fill="black")
    stream = BytesIO()
    image.save(stream, format="PNG")
    image.close()
    scan = target / "scanned.pdf"
    with fitz.open() as doc:
        page = doc.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=stream.getvalue())
        doc.save(scan)
    hybrid = target / "hybrid.pdf"
    with fitz.open(native) as doc, fitz.open(scan) as second:
        doc.insert_pdf(second)
        doc.save(hybrid)
    word = Document()
    for line in LINES:
        word.add_paragraph(line)
    word.save(target / "synthetic.docx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = "\n".join(LINES)
    slides.save(target / "synthetic.pptx")
    faces = create_faces(target / "faces")
    paths = {"text_pdf": native, "scanned_pdf": scan, "hybrid_pdf": hybrid,
             "docx": target / "synthetic.docx", "pptx": target / "synthetic.pptx",
             "face_pptx": faces / "synthetic-face.pptx"}
    write_json(marker, {name: str(path) for name, path in paths.items()})
    return paths


def reopen(path):
    if path.suffix == ".pdf":
        with fitz.open(path) as doc:
            assert len(doc) > 0
    elif path.suffix == ".docx":
        assert Document(path).paragraphs
    else:
        assert len(Presentation(path).slides) > 0


def checked_state(state, *, restored=False):
    assert state["summary"]["failure"] is None, state["summary"]
    if restored:
        # Reload replays saved evidence and regenerates artifacts; extraction and
        # inference are deliberately not rerun and their old timings are absent.
        completed = {s["name"] for s in state["stages"] if s["state"] == "Completed"}
        assert completed == {"artifact_redaction", "artifact_verification"}, state["stages"]
        assert not any(s["state"] == "Failed" for s in state["stages"])
        assert "persistence_load" in state["summary"]["stage_timings_seconds"]
        for flag in ("extraction_completed", "detection_completed", "text_redaction_completed"):
            assert state["summary"][flag]
    else:
        assert all(stage["state"] == "Completed" for stage in state["stages"]), state["stages"]
    assert state["artifact_available"]
    artifact = state["summary"]["artifact"]
    assert artifact["verification_completed"] and artifact["supported_content_verified"], artifact
    for key in ("safe_for_llm", "coverage_verified", "downstream_export_permitted"):
        assert state["summary"][key] is False


def engines(work, paths):
    pipeline = DocumentPipeline(audit=AuditEngine(work / "engine-audits"),
                                artifact_output_dir=work / "engine-artifacts")
    results = {}
    for name, path in paths.items():
        result = pipeline.process(path)
        assert result.failure is None, (name, result.failure)
        extraction = result.extraction_result
        entities = {d.entity_type for d in result.detection_result.detections}
        assert {"EMPLOYEE_ID", "EMAIL_ADDRESS"} <= entities, (name, entities)
        assert "EMP-00001" not in result.redaction_result.full_text
        artifact = result.artifact_result
        assert artifact.reopened and artifact.verification_completed
        assert artifact.supported_content_verified, artifact.sanitized_summary()
        reopen(artifact.path)
        if name == "face_pptx":
            assert artifact.protected_faces == 1 and artifact.face_transformations_verified
        results[name] = {"status": "VERIFIED", "strategy": result.file_intelligence_result.extraction_strategy,
                         "methods": [page.extraction_method for page in extraction.pages],
                         "entities": sorted(entities), "pages": len(extraction.pages),
                         "artifact": artifact.sanitized_summary(), "stages": sorted(result.stage_timings)}
        print(name, "VERIFIED", flush=True)
    assert results["text_pdf"]["methods"] == ["text"]
    assert results["scanned_pdf"]["methods"] == ["ocr"]
    assert results["hybrid_pdf"]["methods"] == ["text", "ocr"]
    write_json(work / "engines.json", results)


def request(client, method, path, **kwargs):
    response = client.request(method, path, **kwargs)
    response.raise_for_status()
    return response


def api(work, paths, client):
    assert request(client, "GET", "/").status_code == 200
    assert request(client, "GET", "/api/config").json()["trusted_local_only"]
    assert client.post("/api/load-audit", json={}).status_code == 403
    records = {}
    for name, path in paths.items():
        state = request(client, "POST", "/api/documents", headers=HEADERS,
                        files={"file": (path.name, path.read_bytes())}).json()
        checked_state(state)
        base = "/api/documents/" + state["document_session_id"]
        found = request(client, "GET", base + "/findings?limit=100").json()
        assert {"EMPLOYEE_ID", "EMAIL_ADDRESS"} <= set(found["entities"])
        employee = next(f for f in found["items"] if f["entity_type"] == "EMPLOYEE_ID")
        detail = request(client, "GET", base + "/findings/" + employee["finding_id"]).json()
        assert detail["candidate"] == "EMP-00001"
        window = request(client, "GET", base + "/source", params={"location": employee["location"]}).json()
        assert "EMP-00001" in window["text"] and window["highlights"]
        artifact = request(client, "GET", base + "/artifact")
        output = work / ("download-" + name + path.suffix)
        output.write_bytes(artifact.content)
        reopen(output)
        records[name] = {"status": "VERIFIED", "stages": state["stages"],
                         "artifact": state["summary"]["artifact"], "download_bytes": len(artifact.content)}
        if name == "docx":
            # Reject then reconfirm a real default-analyzer finding to exercise both paths.
            for action, revision in [("FALSE_POSITIVE", 0), ("CONFIRM_PII", 1)]:
                state = request(client, "POST", base + "/feedback", headers=HEADERS,
                                json={"expected_revision": revision, "finding_id": employee["finding_id"],
                                      "action": action, "reviewer_id": "synthetic-runtime-check"}).json()
                checked_state(state)
                assert state["revision"] == revision + 1
            assert client.post(base + "/feedback", headers=HEADERS,
                               json={"expected_revision": 0, "finding_id": employee["finding_id"],
                                     "action": "CONFIRM_PII"}).status_code == 409
            events = request(client, "GET", base + "/audit").json()["events"]
            assert [event["action"] for event in events] == ["FALSE_POSITIVE", "CONFIRM_PII"]
            store = ROOT / ".private_review_app/audits"
            before = set(store.glob("audit_*.json"))
            saved = request(client, "POST", base + "/save", headers=HEADERS,
                            json={"expected_revision": 2}).json()
            assert saved["saved_revision"] == 2
            new = set(store.glob("audit_*.json")) - before
            assert len(new) == 1, "Expected a unique synthetic audit snapshot"
            snapshot = new.pop()
            payload = json.loads(snapshot.read_text(encoding="utf-8"))
            write_json(work / "restart.json", {"snapshot": str(snapshot),
                       "source": payload["extraction"]["source_path"], "revision": 2})
        print("HTTP", name, "VERIFIED", flush=True)
    write_json(work / "api.json", records)


def reload(work, client):
    saved = json.loads((work / "restart.json").read_text())
    # The current endpoint enumerates sorted snapshot paths; map only our own new file.
    snapshot = Path(saved["snapshot"])
    paths = sorted(snapshot.parent.glob("audit_*.json"))
    refs = request(client, "GET", "/api/saved-audits").json()
    assert len(paths) == len(refs)
    reference = refs[paths.index(snapshot)]["audit_reference"]
    state = request(client, "POST", "/api/load-audit", headers=HEADERS,
                    json={"audit_reference": reference}).json()
    checked_state(state, restored=True)
    assert state["revision"] == saved["revision"]
    base = "/api/documents/" + state["document_session_id"]
    assert len(request(client, "GET", base + "/audit").json()["events"]) == 2
    output = work / "download-after-restart.docx"
    output.write_bytes(request(client, "GET", base + "/artifact").content)
    reopen(output)
    assert "EMP-00001" not in "\n".join(p.text for p in Document(output).paragraphs)
    source = Path(saved["source"])
    assert source.resolve().is_relative_to(ROOT / ".private_review_app/uploads")
    original = source.read_bytes()
    assert original == (work / "fixtures/synthetic.docx").read_bytes()
    try:
        source.write_bytes(original + b"\nsynthetic fingerprint test\n")
        changed = client.post("/api/load-audit", headers=HEADERS, json={"audit_reference": reference})
        assert changed.status_code == 409
    finally:
        source.write_bytes(original)
    restored = request(client, "POST", "/api/load-audit", headers=HEADERS,
                       json={"audit_reference": reference}).json()
    checked_state(restored, restored=True)
    write_json(work / "reload.json", {"status": "VERIFIED", "revision": 2,
               "audit_events": 2, "regenerated_artifact_reopened": True,
               "changed_source_rejected_http": 409, "restored_source_reload": True,
               "coverage_verified": False, "safe_for_llm": False})
    print("Restart, saved listing, reload, event replay and fingerprint rejection VERIFIED", flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("phase", choices=["engines", "api", "reload"])
    parser.add_argument("--work-dir", required=True)
    args = parser.parse_args()
    work = Path(args.work_dir).resolve()
    if not work.is_relative_to(ROOT / ".private_review_app"):
        parser.error("Use a dedicated directory beneath .private_review_app")
    AuditStore(work)._prepare()
    if args.phase == "engines":
        engines(work, fixtures(work))
    else:
        # Fixed loopback URL: this tool never submits documents to external services.
        with httpx.Client(base_url="http://127.0.0.1:8765", timeout=300, trust_env=False) as client:
            if args.phase == "api":
                api(work, fixtures(work), client)
            else:
                reload(work, client)


if __name__ == "__main__":
    main()
